# Run this app with: streamlit run main.py

import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side, Alignment
import os
import datetime

st.title('Mission MRE/CTE Inject Export (3-Line Format)')

# Inputs
descriptions = ["TBM Impact", "Orientation", "CUB", "BUB"]
selected_description = st.selectbox("Choose Inject Description", descriptions)

dtg_input = st.text_input('Enter the DTG for this inject (e.g., 221100ZAPR25):')

if st.button("Generate 3-Line Export"):
    try:
        if not dtg_input:
            st.error("Please enter a DTG before exporting.")
            st.stop()

        # Load inject values from the template
        wb_template = load_workbook("templates/blank_msel_template.xlsx")
        ws_template = wb_template["Sheet1"]

        subject = ws_template["C2"].value
        narrative = ws_template["D1"].value
        audience = ws_template["C3"].value
        from1 = ws_template["E1"].value
        to1 = ws_template["E2"].value
        mode = ws_template["E3"].value
        met = ws_template["F1"].value
        inject_cell = ws_template["F2"].value
        products = ws_template["F3"].value

        # Open the master workbook directly
        master_path = "MSEL.xlsx"
        wb_master = load_workbook(master_path)
        ws_master = wb_master.active
        if ws_master is None or not wb_master.sheetnames:
            ws_master = wb_master.create_sheet("Master MSEL")

        # Prepare inject data block
        row1 = ["002", "", dtg_input, narrative, from1, to1, met]
        row2 = ["", "", subject, "", audience, inject_cell, ""]
        row3 = ["", "", audience, "", mode, "", products]

        # Find next empty inject block starting at row 7
        row_ptr = 7
        while ws_master.cell(row=row_ptr, column=3).value:
            row_ptr += 4

        for i, row_data in enumerate([row1, row2, row3]):
            for j, value in enumerate(row_data, 1):
                ws_master.cell(row=row_ptr + i, column=j, value=value)

        # Apply formatting and merge where needed
        merge_cols = [("A", 3), ("B", 3), ("D", 3)]
        border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        for col, span in merge_cols:
            ws_master.merge_cells(f"{col}{row_ptr}:{col}{row_ptr + span - 1}")

        for row in ws_master.iter_rows(min_row=row_ptr, max_row=row_ptr + 2, min_col=1, max_col=7):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center')

        for col_letter, width in zip("ABCDEFG", [12, 12, 20, 42, 15, 15, 35]):
            ws_master.column_dimensions[col_letter].width = width

        wb_master.save(master_path)
        st.success(f"✅ Inject added directly into {master_path} at rows {row_ptr}-{row_ptr+2}")

    except Exception as e:
        st.error(f"Export failed: {e}")
